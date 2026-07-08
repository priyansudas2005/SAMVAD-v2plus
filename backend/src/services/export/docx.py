"""
docx.py
Microsoft Word DOCX exporter for SAMVAD V2.0.
Uses python-docx to generate proper .docx files with formatting.
"""
import io
from typing import Dict, Any
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .base import BaseExporter, get_export_config, build_export_metadata


def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def _add_bold_paragraph(doc, label, text):
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True
    p.add_run(text)
    return p


class DocxExporter(BaseExporter):

    def export(self, meeting_title: str, date_str: str, segments: list,
               memo: Dict[str, Any] = None,
               intelligence: Dict[str, Any] = None) -> bytes:
        doc = Document()

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        _add_heading(doc, meeting_title, level=0)
        p = doc.add_paragraph()
        p.add_run(f"Date: {date_str}").bold = True

        if memo:
            doc.add_page_break()
            _add_heading(doc, "Executive Summary", level=1)
            doc.add_paragraph(memo.get("summary", "No summary generated."))

            kp = memo.get("key_points", [])
            if kp:
                _add_heading(doc, "Key Highlights", level=2)
                for point in kp:
                    doc.add_paragraph(point, style='List Bullet')

            dp = memo.get("discussion_points", [])
            if dp:
                _add_heading(doc, "Discussion Points", level=2)
                for point in dp:
                    doc.add_paragraph(point, style='List Bullet')

        if intelligence:
            doc.add_page_break()
            _add_heading(doc, "Meeting Intelligence", level=1)

            actions = intelligence.get("action_items", [])
            if actions:
                _add_heading(doc, "Tasks & Action Items", level=2)
                table = doc.add_table(rows=1, cols=5)
                table.style = 'Light Grid Accent 1'
                hdr = table.rows[0].cells
                hdr[0].text = 'Task'
                hdr[1].text = 'Assignee'
                hdr[2].text = 'Priority'
                hdr[3].text = 'Deadline'
                hdr[4].text = 'Status'
                for item in actions:
                    row = table.add_row().cells
                    row[0].text = item.get('task', '')
                    row[1].text = item.get('owner', 'UNKNOWN')
                    row[2].text = item.get('priority', 'MEDIUM')
                    row[3].text = item.get('deadline', 'NONE')
                    row[4].text = item.get('status', 'TODO')

            decisions = intelligence.get("decisions", [])
            if decisions:
                _add_heading(doc, "Key Decisions", level=2)
                for dec in decisions:
                    text = dec.get("text") if isinstance(dec, dict) else str(dec)
                    dec_type = dec.get("type", "FINAL") if isinstance(dec, dict) else ""
                    speakers = dec.get("supporting_speakers", []) if isinstance(dec, dict) else []
                    suffix = f" [{dec_type}]" if dec_type else ""
                    if speakers:
                        suffix += f" (by {', '.join(speakers[:3])})"
                    doc.add_paragraph(f"• {text}{suffix}")

            risks = intelligence.get("risks", [])
            if risks:
                _add_heading(doc, "Risks Identified", level=2)
                for r in risks:
                    text = r.get("text") if isinstance(r, dict) else str(r)
                    doc.add_paragraph(f"Risk: {text}", style='List Bullet')

        doc.add_page_break()
        _add_heading(doc, "Detailed Transcript", level=1)
        for seg in segments:
            speaker = seg.get("speaker_label", "UNKNOWN")
            ts = f"[{seg.get('start', '00:00')} - {seg.get('end', '00:00')}]"
            p = doc.add_paragraph()
            run = p.add_run(f"{ts} ")
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 116, 139)
            run2 = p.add_run(f"{speaker}: ")
            run2.bold = True
            run2.font.size = Pt(10)
            p.add_run(seg.get('text', ''))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()
